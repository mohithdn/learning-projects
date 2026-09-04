"""
Project 20 (CAPSTONE): Autonomous Research & Executive Slide Deck Generator
------------------------------------------------------------------------------
Goal: Enter an industry/topic -> Gemini conducts live grounded web research
-> structures findings into a 5-slide JSON outline -> app renders 5 preview
cards -> user downloads an actual .pptx file built from that outline.

Run with:
    python -m streamlit run app_p20.py

Requires:
    pip install streamlit google-genai python-pptx --break-system-packages
"""

import io
import json

import streamlit as st
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# ---------------------------------------------------------
# Page Config
# ---------------------------------------------------------
st.set_page_config(
    page_title="Executive Slide Deck Generator",
    page_icon="📊",
    layout="centered",
)

st.title("📊 Autonomous Research & Slide Deck Generator")
st.caption("Enter an industry or topic — the agent researches it live on the web, structures a 5-slide outline, and exports a downloadable PowerPoint.")

GEN_MODEL = "gemini-3.7-flash"

# ---------------------------------------------------------
# Sidebar: API Key
# ---------------------------------------------------------
with st.sidebar:
    st.header("⚙️ Configuration")
    api_key = st.text_input("Gemini API Key", type="password", placeholder="Paste your key here")
    st.markdown("---")
    st.markdown(
        "**Pipeline:**\n"
        "1️⃣ Live grounded research (Google Search tool)\n"
        "2️⃣ Structure into 5-slide JSON\n"
        "3️⃣ Preview cards\n"
        "4️⃣ Export as `.pptx`"
    )

# ---------------------------------------------------------
# Session State
# ---------------------------------------------------------
if "slide_data" not in st.session_state:
    st.session_state.slide_data = None
if "research_notes" not in st.session_state:
    st.session_state.research_notes = ""

# ---------------------------------------------------------
# Main Input
# ---------------------------------------------------------
topic = st.text_input("Industry or topic to research:", placeholder="e.g. The state of solid-state EV batteries in 2026")

generate_clicked = st.button("🚀 Research & Generate Deck", type="primary")

# ---------------------------------------------------------
# STEP 1: Live grounded research
# ---------------------------------------------------------
if generate_clicked:
    if not api_key:
        st.warning("⚠️ Please enter your Gemini API key in the sidebar.")
    elif not topic.strip():
        st.warning("⚠️ Please enter a topic first.")
    else:
        try:
            client = genai.Client(api_key=api_key)

            with st.spinner("Step 1/2 · Conducting live web research..."):
                grounding_tool = types.Tool(google_search=types.GoogleSearch())
                research_config = types.GenerateContentConfig(tools=[grounding_tool])

                research_prompt = (
                    f"Research the current state of '{topic}'. Gather the most "
                    f"important, up-to-date facts, trends, key players, and figures "
                    f"a business executive would want to know. Write your findings "
                    f"as clear factual notes."
                )

                research_response = client.models.generate_content(
                    model=GEN_MODEL,
                    contents=research_prompt,
                    config=research_config,
                )
                st.session_state.research_notes = research_response.text

            # -----------------------------------------------------
            # STEP 2: Structure research into a strict 5-slide JSON outline
            # -----------------------------------------------------
            with st.spinner("Step 2/2 · Structuring into a 5-slide outline..."):
                structuring_prompt = f"""
                Convert the research notes below into a 5-slide executive
                presentation outline.

                Return ONLY valid JSON, no markdown fences, no preamble, in this
                exact schema:
                {{
                  "deck_title": "string",
                  "slides": [
                    {{
                      "slide_number": 1,
                      "title": "string (short slide title)",
                      "bullets": ["string", "string", "string"]
                    }}
                  ]
                }}

                Requirements:
                - Exactly 5 slides.
                - Slide 1 should be a title/overview slide (1-2 bullets max).
                - Slides 2-4 should each cover a distinct key theme with 3-4 concise bullets.
                - Slide 5 should be a "Key Takeaways" or "Outlook" summary slide.
                - Bullets must be concise (under 15 words each), based strictly on
                  the research notes — do not invent facts not present below.

                Research notes:
                \"\"\"
                {st.session_state.research_notes}
                \"\"\"
                """

                structure_config = types.GenerateContentConfig(
                    response_mime_type="application/json",
                )

                structure_response = client.models.generate_content(
                    model=GEN_MODEL,
                    contents=structuring_prompt,
                    config=structure_config,
                )

                st.session_state.slide_data = json.loads(structure_response.text)

            st.success("✅ Research complete and deck outline generated.")

        except json.JSONDecodeError:
            st.error("❌ The model did not return valid JSON for the outline. Try clicking Generate again.")
        except Exception as e:
            st.error(f"❌ Something went wrong: {e}")

# ---------------------------------------------------------
# STEP 3: Preview cards
# ---------------------------------------------------------
if st.session_state.slide_data:
    deck = st.session_state.slide_data
    st.markdown(f"## 🗂️ {deck.get('deck_title', 'Untitled Deck')}")

    with st.expander("📝 Raw research notes used"):
        st.write(st.session_state.research_notes)

    for slide in deck.get("slides", []):
        with st.container(border=True):
            st.markdown(f"**Slide {slide['slide_number']} · {slide['title']}**")
            for bullet in slide["bullets"]:
                st.markdown(f"- {bullet}")

    st.markdown("---")

    # -----------------------------------------------------
    # STEP 4: Build the actual .pptx file from the JSON outline
    # -----------------------------------------------------
    def build_pptx(deck: dict) -> bytes:
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        title_layout = prs.slide_layouts[0]
        content_layout = prs.slide_layouts[1]

        # --- Title slide ---
        title_slide = prs.slides.add_slide(title_layout)
        title_slide.shapes.title.text = deck.get("deck_title", "Untitled Deck")
        if len(title_slide.placeholders) > 1:
            first_slide_bullets = deck["slides"][0]["bullets"] if deck.get("slides") else []
            title_slide.placeholders[1].text = "\n".join(first_slide_bullets)

        # --- Content slides (slides 2 onward in the JSON become the body) ---
        for slide_info in deck.get("slides", [])[1:]:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = slide_info["title"]

            body = slide.placeholders[1].text_frame
            body.clear()
            for i, bullet in enumerate(slide_info["bullets"]):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = bullet
                p.font.size = Pt(20)

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    pptx_bytes = build_pptx(deck)

    st.download_button(
        label="⬇️ Download Slide Deck (.pptx)",
        data=pptx_bytes,
        file_name=f"{deck.get('deck_title', 'deck').replace(' ', '_')[:50]}.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        type="primary",
    )

# ---------------------------------------------------------
# Footer
# ---------------------------------------------------------
st.markdown("---")
st.caption("Project 20 · CAPSTONE · Autonomous Research & Executive Slide Deck Generator · Built with Streamlit + Gemini")